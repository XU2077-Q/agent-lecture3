# -*- coding: utf-8 -*-
"""
生成两份 Word 文档：
  1. 第三讲-讲解.docx       —— 课件逐节讲解 + 术语解释
  2. Agent功能说明.docx      —— 本仓库 Agent 的主要功能说明
输出到桌面与本仓库 docs/ 各一份。
运行：python gen_docs.py
"""
import os
import shutil

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

DESKTOP = os.path.join(os.path.expanduser("~"), "Desktop")
DOCS_DIR = os.path.join(os.path.dirname(os.path.abspath(__file__)), "docs")


# ---------- 排版辅助 ----------

def set_cn_fonts(doc):
    """让中文用微软雅黑渲染，避免默认字体发虚。"""
    for name in ("Normal", "Heading 1", "Heading 2", "Heading 3", "Title"):
        try:
            st = doc.styles[name]
        except KeyError:
            continue
        st.font.name = "Calibri"
        rpr = st.element.get_or_add_rPr()
        rpr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    doc.styles["Normal"].font.size = Pt(11)


def h(doc, text, level=1):
    doc.add_heading(text, level)


def p(doc, text, bold=False):
    para = doc.add_paragraph()
    run = para.add_run(text)
    run.bold = bold
    return para


def bullets(doc, items):
    for it in items:
        doc.add_paragraph(it, style="List Bullet")


def code(doc, text):
    for line in text.strip("\n").split("\n"):
        para = doc.add_paragraph()
        para.paragraph_format.space_after = Pt(0)
        para.paragraph_format.left_indent = Pt(18)
        run = para.add_run(line)
        run.font.name = "Consolas"
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(0x1F, 0x3F, 0x7F)


def table(doc, headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    for i, head in enumerate(headers):
        cell = t.rows[0].cells[i]
        cell.text = head
        for run in cell.paragraphs[0].runs:
            run.bold = True
    for row in rows:
        cells = t.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = val
    return t


def save(doc, filename):
    os.makedirs(DOCS_DIR, exist_ok=True)
    local = os.path.join(DOCS_DIR, filename)
    doc.save(local)
    shutil.copy(local, os.path.join(DESKTOP, filename))
    print("已生成:", local, "和桌面副本")


# ================= 文档一：第三讲-讲解.docx =================

def build_lecture_doc():
    doc = Document()
    set_cn_fonts(doc)

    title = doc.add_heading("《Agent 外部通讯 · 第三讲》讲解", 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p(doc, "课程：中山大学《软工新生研讨实践》　　整理日期：2026-10-14")
    p(doc, "配套代码仓库：agent-lecture3（echo_bot.py / llm.py / wechat_bot.py / feishu_bot.py）")

    # ---------- 一 ----------
    h(doc, "一、本讲任务与整体思路", 1)
    p(doc, "本讲只有一个核心任务：让一个聊天软件里的“机器人”替你收发消息，并把消息交给大模型处理。"
           "把这句话拆开，就是一个永远转着的循环：")
    bullets(doc, [
        "收消息——机器人账号收到用户发来的文字；",
        "调大模型——把这段文字发给 AI 模型，拿到回答；",
        "回消息——把回答从机器人账号发回给用户。",
    ])
    p(doc, "课程先在命令行里把这个循环跑通（复读机），再分别接入两个真实平台："
           "微信（走 iLink 官方机器人接口）和飞书（走自建应用 + 长连接）。"
           "两个平台“收、调、回”的主干完全一样，差别只在身份怎么来、消息怎么收、回复用什么载体。")

    # ---------- 二 ----------
    h(doc, "二、微信机器人（iLink）：十步走", 1)

    h(doc, "2.0 第零步：命令行复读机", 2)
    p(doc, "先用 6 行 Python 体会“循环收、循环回”，这是后面一切的地基（完整文件见仓库 echo_bot.py）：")
    code(doc, '''
while True:
    text = input("你: ")
    if text == "quit":
        break
    print("机器人: " + text)''')

    h(doc, "2.1 微信消息流转总览（四步）", 2)
    bullets(doc, [
        "① 用户在微信 App 里给机器人发消息，消息先到腾讯的 iLink 服务器暂存；",
        "② 我们的程序主动调用 getupdates 把消息取回来（拉的模式）；",
        "③ 程序把消息文字交给大模型，得到回答；",
        "④ 程序调用 sendmessage，让 iLink 把回答投递回用户的聊天窗口。",
    ])
    p(doc, "要点：微信端是我们的程序“主动去问服务器有没有新消息”，飞书端则相反，是服务器主动推给我们。")

    h(doc, "2.2 第 1~2 步：取登录二维码", 2)
    p(doc, "调用 get_bot_qrcode（参数 bot_type=3）拿到二维码内容，"
           "再用 Python 的 qrcode 库把内容画成图片 qr_wechat.png：")
    code(doc, '''
r = requests.get("https://ilinkai.weixin.qq.com/ilink/bot/get_bot_qrcode",
                 params={"bot_type": 3})
data = r.json()
qrcode.make(data["qrcode_img_content"]).save("qr_wechat.png")''')

    h(doc, "2.3 第 3 步：扫码登录拿 bot_token", 2)
    p(doc, "用户用微信扫码后，程序每隔约 2 秒轮询一次 get_qrcode_status，"
           "状态从 wait（未扫）→ scaned（已扫待确认）→ confirmed（已确认）。"
           "确认后响应里带着 bot_token——这就是机器人账号的“钥匙”，写入 wx_token.txt 以后复用：")
    code(doc, '''
if status == "confirmed":
    token = s["bot_token"]
    open("wx_token.txt", "w").write(token)''')

    h(doc, "2.4 第 4 步：组装请求头", 2)
    p(doc, "之后每个请求都要带四个头，向 iLink 证明“我是那个扫码登录过的机器人”：")
    code(doc, '''
{"Content-Type": "application/json",
 "AuthorizationType": "ilink_bot_token",
 "Authorization": "Bearer " + token,
 "X-WECHAT-UIN": uin}    # 一个数字的 base64，课件示例可直接固定（如 10086）''')

    h(doc, "2.5 第 5 步：长轮询收消息（getupdates）", 2)
    p(doc, "POST getupdates 时带上游标 get_updates_buf。服务器不会立刻返回，而是最多“挂住”约 35 秒："
           "期间有新消息就立即返回；没有就空手而归，程序带着原游标再问一轮。"
           "游标相当于书签——记住“我收到哪了”，断线重连也不丢消息。"
           "原始返回包顺手写进 last_msg.json，调试时直接翻开看：")
    code(doc, '''
r = requests.post(BASE + "/ilink/bot/getupdates", headers=get_headers(token),
                  json={"get_updates_buf": cursor}, timeout=45)
cursor = r.json().get("get_updates_buf", cursor)''')

    h(doc, "2.6 第 6 步：解析消息", 2)
    p(doc, "返回包是消息列表。只挑两类字段：message_type==1 表示“用户发来的消息”（区别于系统通知等）；"
           "item_list 里 type==1 的那一项才是文字（2 图片、3 语音、4 文件、5 视频）。"
           "取出三样东西：发件人 from_user_id、会话令牌 context_token、文字内容，"
           "前两者原样记下，回复时要用：")
    code(doc, '''
if msg["message_type"] == 1:
    uid = msg["from_user_id"]
    ctx = msg["context_token"]
    text = msg["item_list"][0]["text_item"]["text"]''')

    h(doc, "2.7 第 7 步：回消息（sendmessage）", 2)
    p(doc, "POST sendmessage。最关键的是 context_token 必须原样带回——它相当于“回话地址”，"
           "决定这条回复落进哪个聊天窗口。client_id 用 uuid 保证全局唯一，防止网络重试造成重复发送：")
    code(doc, '''
{"msg": {"from_user_id": "", "to_user_id": uid,
         "client_id": "claw-" + uuid.uuid4().hex[:12],
         "message_type": 2, "message_state": 2, "context_token": ctx,
         "item_list": [{"type": 1, "text_item": {"text": reply}}],
         "base_info": {"channel_version": "1.0.3"}}}''')

    h(doc, "2.8 第 8~9 步：串成循环、接入大模型", 2)
    p(doc, "把“收 → 解析 → 回”套进 while True 主循环，复读机就联网了；"
           "再在中间插入一次大模型调用，复读机就升级成 AI 机器人。"
           "大模型接口用 OpenAI 兼容格式（/chat/completions），换服务商只改地址和密钥：")
    code(doc, '''
def ask_llm(text):
    r = requests.post(LLM_BASE + "/chat/completions",
        headers={"Authorization": "Bearer " + LLM_KEY},
        json={"model": "gpt-4o-mini",
              "messages": [{"role": "user", "content": text}]})
    return r.json()["choices"][0]["message"]["content"]''')

    h(doc, "2.9 第 10 步：人设（system 角色）", 2)
    p(doc, "messages 里可以塞一条 role=system 的消息，它不来自用户，而是“导演给演员的剧本”："
           "规定模型是谁、用什么口吻。课件例子是让模型扮周杰伦粉丝：")
    code(doc, '''
messages = [{"role": "system", "content": "你是周杰伦的粉丝，回答都向着偶像"},
            {"role": "user", "content": "你觉得周杰伦怎么样"}]''')
    p(doc, "微信端是“整段问、整段答”：拿到完整回答后一次发出，微信接口本身不支持逐字显示。")

    # ---------- 三 ----------
    h(doc, "三、飞书机器人：自建应用 + 长连接", 1)

    h(doc, "3.1 创建自建应用", 2)
    bullets(doc, [
        "在 open.feishu.cn/app 创建“自建应用”，拿到 App ID 和 App Secret；",
        "App Secret 就是密码：不发给任何人、不截图外传、不进代码仓库；",
        "开启“机器人”能力，机器人才有资格收发消息。",
    ])

    h(doc, "3.2 权限、事件与发布", 2)
    bullets(doc, [
        "权限管理开通三个：im:message（收发消息）、im:message:send_as_bot（以机器人身份发）、cardkit:card:write（写卡片）；",
        "事件订阅选择“使用长连接接收事件”——不用买服务器、不用填公网地址，这是飞书对新手最友好的一点；",
        "添加事件“接收消息 im.message.receive_v1”；",
        "创建版本并发布后机器人才生效；以后每改一次配置都要重新发布。",
    ])

    h(doc, "3.3 长连接收消息", 2)
    p(doc, "飞书端不需要轮询：官方 SDK（lark-oapi）用 WebSocket 与服务器保持一条长连接，"
           "有消息时服务器顺着连接直接推过来，毫秒级到达。注册一个回调函数即可：")
    code(doc, '''
handler = (lark.EventDispatcherHandler.builder("", "")
           .register_p2_im_message_receive_v1(on_msg)
           .build())
lark.ws.Client(APP_ID, APP_SECRET, event_handler=handler).start()''')
    p(doc, "回调里第一守则：机器人自己发的消息（sender_type 为 app）直接忽略，"
           "否则机器人会对自己的回答再回答，陷入死循环。")

    h(doc, "3.4 消息卡片与 Markdown", 2)
    p(doc, "飞书的回复可以不是纯文字，而是一张“消息卡片”（msg_type=interactive）。"
           "卡片里的 markdown 组件支持常用语法：加粗 **文字**、斜体 *文字*、列表、链接 [文字](网址)、分割线 ---。"
           "卡片发出后还能修改内容——这正是逐字显示的前提。")

    h(doc, "3.5 流式输出（CardKit 四步）", 2)
    p(doc, "想让回答像 ChatGPT 那样一个字一个字冒出来，用 CardKit 流式更新，共四步：")
    bullets(doc, [
        "① 建卡片实体，配置里开启 streaming_mode（流式模式）；",
        "② 把卡片作为 interactive 消息发到会话，用户先看到一张空卡片；",
        "③ 每隔约 0.3 秒“全量替换”一次 markdown 内容（每次发的都是截至目前的完整文本），"
        "请求里的 sequence 序号必须严格递增，更新太密会被限流；",
        "④ 内容发完，关闭流式模式，卡片定格。",
    ])
    p(doc, "稳健做法：卡片流程任何一步失败，就降级成普通文本消息，保证用户一定收得到回复。"
           "本仓库 feishu_bot.py 就是这个思路。")

    # ---------- 四 ----------
    h(doc, "四、微信 vs 飞书：一张对比表", 1)
    table(doc,
          ["维度", "微信（iLink）", "飞书（自建应用）"],
          [["身份获取", "扫码绑定，拿到 bot_token", "创建自建应用，拿 App ID / App Secret"],
           ["收消息", "主动长轮询 getupdates（约 35 秒）", "WebSocket 长连接，服务器推送"],
           ["回复载体", "纯文字", "可修改的消息卡片（支持 Markdown）"],
           ["逐字显示", "不支持，整段发出", "支持，CardKit 流式更新"],
           ["相同点", "收消息 → 调大模型 → 回消息", "收消息 → 调大模型 → 回消息"]])

    # ---------- 五 ----------
    h(doc, "五、延伸内容", 1)
    bullets(doc, [
        "接收媒体文件：item_list 里 type 2 图片、3 语音、4 文件、5 视频；"
        "语音消息会附带转文字结果；拿到媒体后先存本地，再交给大模型处理。",
        "其他可接的通道：企业微信客服消息接口、钉钉（Stream 长连接 / Webhook）、邮件（SMTP 发 / IMAP 收）。",
        "万变不离其宗：任何通道都是“收消息 → 调大模型 → 回消息”这三段，只是收发的协议不同。",
    ])

    # ---------- 六 ----------
    h(doc, "六、课后任务", 1)
    bullets(doc, [
        "把你的微信或你的飞书，接入你的 AI 模型；",
        "截个图提交；",
        "交了就有分，重在实践，不要抄，不要卷；",
        "之前的任务接受补交，但要尽快，以免后面跟不上。",
    ])
    p(doc, "本仓库 wechat_bot.py（微信）与 feishu_bot.py（飞书）即为该任务的完整实现，两个平台任选其一即可交差。")

    # ---------- 七 ----------
    h(doc, "七、术语解释", 1)
    p(doc, "按课件出现顺序整理，供复习速查：")
    table(doc,
          ["术语", "解释"],
          [["Agent（智能体）", "能感知输入、调用模型/工具、产出行动的程序。本讲的 Agent 就是一个住在聊天软件里的对话机器人。"],
           ["iLink", "腾讯微信团队开放的官方机器人接口（ilinkai.weixin.qq.com），扫码即可让程序以机器人身份收发微信消息。"],
           ["bot_token", "扫码确认后 iLink 发给的令牌，之后所有接口调用都靠它证明身份。等于机器人账号的钥匙，须保密。"],
           ["get_bot_qrcode / get_qrcode_status", "iLink 的两个登录接口：前者取二维码内容，后者轮询扫码状态（wait→scaned→confirmed）。"],
           ["长轮询（Long Polling）", "客户端发请求后服务器不立即返回，而是挂住一段时间（iLink 约 35 秒），有消息立刻回、没消息超时回空，客户端马上再发下一轮。兼顾实时性与实现简单。"],
           ["getupdates", "iLink 的收消息接口，采用长轮询方式调用。"],
           ["cursor / get_updates_buf", "游标，相当于书签：记录“消息收到哪了”。每次请求带上次返回的游标，断线重连也不重复、不丢消息。"],
           ["message_type / item_list", "iLink 消息结构。message_type==1 表示用户消息；item_list 是消息内容列表，其中 type 1 文字、2 图片、3 语音、4 文件、5 视频。"],
           ["context_token", "会话令牌，相当于“回话地址”。回复时原样带回，iLink 才知道把消息投递到哪个聊天窗口。"],
           ["from_user_id / to_user_id", "发件人 / 收件人标识。回复时把收到的 from_user_id 填进 to_user_id，就回给了原发信人。"],
           ["client_id", "客户端为每条消息生成的唯一编号（课件用 claw- 前缀加 uuid 前 12 位），用于防重：网络重试时服务器能识别同一条消息。"],
           ["X-WECHAT-UIN", "iLink 请求头之一，一个数字的 base64 编码，标识微信用户域。课件说明可固定取值（如 10086）。"],
           ["base64", "一种把任意字节编码成可打印字符的方案，不是加密。Python 用 base64.b64encode。"],
           ["LLM / 大模型", "Large Language Model，大语言模型，如 GPT、DeepSeek、通义千问。本讲用它生成机器人回复。"],
           ["OpenAI 兼容接口", "业界事实标准：POST {BASE}/chat/completions，带 model 与 messages 数组。换国产模型只需改 BASE、模型名和密钥。"],
           ["messages / role", "发给大模型的对话数组。role 有三种常见取值：system（系统设定）、user（用户输入）、assistant（模型回答）。"],
           ["system 角色 / 人设", "system 消息是“导演给演员的剧本”，规定模型的身份与口吻，如“你是周杰伦的粉丝”。也称人设（Persona）。"],
           ["API Key / Bearer Token", "调用大模型接口的密钥，放在 Authorization: Bearer {key} 请求头里。泄露等于账号被盗用，绝不提交进仓库。"],
           ["自建应用", "在飞书开放平台由你自己创建、只服务于你自己组织的应用。与之相对的是上架应用市场的“商店应用”。"],
           ["App ID / App Secret", "自建应用的账号与密码。App Secret 绝不外传；泄露后应立即在后台重置。"],
           ["tenant_access_token", "飞书应用级访问令牌，用 App ID/Secret 换取，约 2 小时有效，调用发消息等 OpenAPI 时放在 Bearer 头里。"],
           ["WebSocket 长连接", "一条建立后长期保持的双向连接。飞书 SDK 靠它接收事件推送，免公网 IP、免域名备案，实时性优于轮询。"],
           ["事件订阅 / im.message.receive_v1", "告诉飞书“哪些事要通知我”。im.message.receive_v1 即“收到单聊消息”事件，v1 是事件版本号。"],
           ["消息卡片 / interactive", "飞书的富文本消息类型（msg_type=interactive），可排版、带按钮，且发出后内容可修改。"],
           ["Markdown", "轻量排版语法：**加粗**、*斜体*、- 列表、1. 有序列表、[文字](网址) 链接、--- 分割线。飞书卡片 markdown 组件直接支持。"],
           ["CardKit", "飞书卡片平台接口，提供卡片实体的创建、更新能力，是流式输出的实现基础。需 cardkit:card:write 权限。"],
           ["流式输出 / streaming_mode", "让回答逐字出现的效果。原理是反复“全量替换”卡片里的文本——每次发送的都是截至当前的完整内容，不是增量片段。"],
           ["sequence", "卡片更新请求里严格递增的序号，防止乱序更新把新内容覆盖成旧的。"],
           ["限流（Rate Limit）", "平台对调用频率的限制。课件经验：卡片更新间隔约 0.3 秒，太密会被拒。"],
           ["降级（Fallback）", "高级通道失败时退回简单通道。本仓库：流式卡片失败 → 退回纯文本消息，保证回复必达。"],
           ["Webhook", "“有事发 POST 给你”的通知方式，需要你有公网可达的服务器地址。钉钉机器人常用；飞书长连接方案则免公网地址。"],
           ["SMTP / IMAP", "邮件两大协议：SMTP 负责发信，IMAP 负责收信。把机器人接到邮箱就是“邮件通道”。"],
           ["chat_id", "飞书会话（聊天窗口）的唯一编号，发消息时用它指定发到哪个会话。"]])

    p(doc, "")
    p(doc, "附：课件末两页为“AI+教育”创新应用技能大赛宣传（aiedu.caet.org.cn，中山大学赛区）"
           "与课程匿名反馈二维码，与技术内容无关，从略。")

    save(doc, "第三讲-讲解.docx")


# ================= 文档二：Agent功能说明.docx =================

def build_agent_doc():
    doc = Document()
    set_cn_fonts(doc)

    title = doc.add_heading("ClawBot：双平台聊天 Agent 功能说明", 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p(doc, "对应课后任务：“把你的微信或你的飞书，接入你的 AI 模型。”　　整理日期：2026-10-14")
    p(doc, "仓库：agent-lecture3　　语言：Python 3　　依赖：requests / qrcode / lark-oapi")

    h(doc, "一、Agent 是什么", 1)
    p(doc, "本 Agent 是一个住在聊天软件里的 AI 对话机器人。它反复执行同一个三段式循环："
           "收消息 → 调大模型 → 回消息。用户在微信或飞书里像跟真人聊天一样发文字，"
           "机器人把文字交给大语言模型，再把模型回答发回来。"
           "同一个“大脑”（llm.py）服务两个“身体”（微信端、飞书端），互不影响，可同时在线。")

    h(doc, "二、总体架构", 1)
    code(doc, '''
        微信用户 ──► iLink 服务器 ──► wechat_bot.py ─┐
                                                    ├─► llm.py ──► 大模型（OpenAI 兼容接口）
        飞书用户 ──► 飞书开放平台 ──► feishu_bot.py ─┘          │
                     ◄── 流式卡片 / 纯文本 ◄────────────────────┘''')
    bullets(doc, [
        "llm.py：唯一的大模型出口。OpenAI 兼容接口，换服务商只改环境变量；system 人设集中在此。",
        "wechat_bot.py：微信通道。iLink 扫码登录，长轮询收消息，纯文字回复。",
        "feishu_bot.py：飞书通道。WebSocket 长连接收消息，CardKit 流式卡片回复，失败自动降级纯文本。",
        "echo_bot.py：第零步命令行复读机，用于不联网验证最小循环。",
    ])

    h(doc, "三、主要功能清单", 1)

    h(doc, "3.1 大模型对话（llm.py）", 2)
    bullets(doc, [
        "ask_llm(text, system)：整段提问、整段返回，一次 HTTP 调用完成；",
        "OpenAI 兼容：默认 https://api.openai.com/v1 + gpt-4o-mini，"
        "用环境变量 LLM_BASE / LLM_MODEL 可切到 DeepSeek、通义、Kimi、火山方舟等任意兼容服务；",
        "密钥双通道：环境变量 LLM_KEY 优先，其次读取本地 llm_key.txt（已在 .gitignore，绝不入库）；",
        "人设（Persona）：通过 system 角色消息设定机器人性格，一行字符串即可换“演员剧本”；",
        "故障兜底：两个机器人端都捕获大模型异常，模型挂了主循环不死，用户收到错误提示而非沉默。",
    ])

    h(doc, "3.2 微信端（wechat_bot.py）", 2)
    bullets(doc, [
        "扫码登录：get_bot_qrcode 取二维码 → qrcode 库生成 qr_wechat.png → 轮询状态 "
        "wait→scaned→confirmed → bot_token 写入 wx_token.txt，之后启动免扫码；",
        "长轮询收消息：getupdates 服务器挂住约 35 秒，游标 get_updates_buf 当书签，断线续传不丢消息；"
        "超时与网络抖动自动进入下一轮，主循环永不退出；",
        "消息解析：只响应 message_type==1 的用户消息中 type==1 的文字项，"
        "自动跳过系统通知与图片/语音/文件/视频；原始包落盘 last_msg.json 便于调试；",
        "可靠回复：context_token 原样带回确保回复落到正确会话；"
        "client_id 用 uuid 生成，防网络重试导致的重复发送；",
        "请求头合规：Authorization Bearer + AuthorizationType + X-WECHAT-UIN（base64）+ channel_version 1.0.3。",
    ])

    h(doc, "3.3 飞书端（feishu_bot.py）", 2)
    bullets(doc, [
        "零公网部署：lark-oapi SDK 建立 WebSocket 长连接，事件由服务器推送，"
        "不需要买服务器、不需要域名备案；",
        "防自回环：sender_type 为 app 的消息（机器人自己发的）直接忽略，杜绝机器人对自己回答再回答的死循环；",
        "流式卡片回复：CardKit 四步——建卡开 streaming_mode → 发卡到会话 → 每 0.3 秒全量替换文本"
        "（sequence 严格递增，防限流防乱序）→ 关闭流式定格，回答像真人打字一样逐字出现；",
        "Markdown 排版：卡片支持加粗、斜体、列表、链接、分割线，回答结构化可读；",
        "自动降级：卡片流程任何一步失败，立即退回普通文本消息，保证回复必达；",
        "非文字消息应答：收到图片等非文字消息时礼貌提示“目前只看得懂文字消息”，不静默。",
    ])

    h(doc, "3.4 命令行复读机（echo_bot.py）", 2)
    bullets(doc, [
        "6 行代码的最小循环：input 收、print 回，输入 quit 退出；",
        "用途：不联网验证 Python 环境与“收→回”心智模型，是两位“真身”的教学起点。",
    ])

    h(doc, "四、安全设计", 1)
    bullets(doc, [
        "密钥零入库：llm_key.txt、wx_token.txt、fs_app_id.txt、fs_app_secret.txt、二维码与消息留底文件全部列入 .gitignore；",
        "环境变量优先：生产用法推荐 LLM_KEY / FS_APP_ID / FS_APP_SECRET 注入，磁盘上不留明文；",
        "App Secret 即密码：文档与 README 均显式提醒不外发、泄露即重置；",
        "最小权限：飞书端仅申请课件要求的三个权限（im:message、im:message:send_as_bot、cardkit:card:write）。",
    ])

    h(doc, "五、运行方式速查", 1)
    table(doc,
          ["步骤", "命令 / 操作"],
          [["装依赖", "pip install -r requirements.txt"],
           ["配模型密钥", "set/export LLM_KEY=sk-...（或新建 llm_key.txt）"],
           ["命令行复读机", "python echo_bot.py"],
           ["微信机器人", "python wechat_bot.py（首次扫 qr_wechat.png）"],
           ["飞书机器人", "开放平台建应用、配权限事件并发布后：python feishu_bot.py"],
           ["生成文档", "python gen_docs.py（输出 docs/ 与桌面各一份）"]])

    h(doc, "六、文件清单", 1)
    table(doc,
          ["文件", "功能"],
          [["echo_bot.py", "命令行复读机（教学第零步）"],
           ["llm.py", "大模型调用封装 + 人设 + 密钥加载"],
           ["wechat_bot.py", "微信 iLink 通道：登录、长轮询、解析、回复"],
           ["feishu_bot.py", "飞书通道：长连接收事件、流式卡片、降级纯文本"],
           ["gen_docs.py", "生成本文档与《第三讲-讲解》"],
           ["requirements.txt", "依赖清单"],
           ["README.md", "仓库总说明与上手步骤"],
           ["docs/第三讲-讲解.docx", "课件逐节讲解 + 术语解释"],
           ["docs/Agent功能说明.docx", "本文档"]])

    save(doc, "Agent功能说明.docx")


if __name__ == "__main__":
    build_lecture_doc()
    build_agent_doc()
